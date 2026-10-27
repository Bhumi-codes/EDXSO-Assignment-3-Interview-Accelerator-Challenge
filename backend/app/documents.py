from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from pypdf import PdfReader
from docx import Document
from fastapi import HTTPException

MAX_BYTES = 5 * 1024 * 1024

def extract_document(data: bytes, filename: str, limit: int = 14000) -> str:
    if not data:
        raise HTTPException(422,'The file is empty. Paste text or choose another file.')
    if len(data) > MAX_BYTES:
        raise HTTPException(413,'Choose a file smaller than 5 MB.')
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == '.txt':
            output = data.decode('utf-8-sig')
        elif suffix == '.pdf':
            if not data.startswith(b'%PDF-'):
                raise ValueError()
            pdf = PdfReader(BytesIO(data))
            if pdf.is_encrypted or len(pdf.pages) > 80:
                raise HTTPException(422,'Use an unencrypted PDF with at most 80 pages.')
            parts = []
            for page in pdf.pages:
                parts.append(page.extract_text() or '')
                if sum(map(len,parts)) > limit:
                    raise HTTPException(422,f'Extracted text exceeds {limit:,} characters. Paste only the relevant content.')
            output = '\n'.join(parts)
        elif suffix == '.docx':
            # Guard expanded size before python-docx opens the ZIP container.
            with ZipFile(BytesIO(data)) as archive:
                if len(archive.infolist()) > 500 or sum(i.file_size for i in archive.infolist()) > 20*1024*1024:
                    raise HTTPException(422,'This DOCX expands beyond the supported size.')
                if 'word/document.xml' not in archive.namelist():
                    raise ValueError()
            doc = Document(BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            parts.extend(' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
            output = '\n'.join(parts)
        else:
            raise HTTPException(415,'Supported formats: PDF, DOCX and UTF-8 TXT.')
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422,'Unable to read this file. Use an unencrypted text-based PDF, DOCX or UTF-8 TXT, or paste the text.') from None
    output = output.replace('\x00','').strip()
    if len(output) < 20:
        raise HTTPException(422,'Not enough readable text. Scanned PDFs need a text-based export or pasted text; OCR is not available.')
    if len(output) > limit:
        raise HTTPException(422,f'Extracted text exceeds {limit:,} characters. Paste only the relevant content.')
    return output
